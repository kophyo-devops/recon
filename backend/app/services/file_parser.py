"""Parse uploaded Excel / CSV into normalized row dicts (data only, no file storage)."""

from __future__ import annotations

import csv
import io
import re
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".xls"}

APP_ID_HEADERS = {
    "application_id",
    "applicationid",
    "app_id",
    "ceir_application_id",
}
REQUEST_HEADERS = {
    "request_id",
    "requestid",
}
STATUS_HEADERS = {
    "status",
    "api_status",
    "apistatus",
    "payment_status",
    "application_status",
    "applicationstatus",
}
# Priority order for a single amount column (first non-empty wins).
AMOUNT_HEADER_PRIORITY = (
    "tax_amount_mmk",
    "tax_amount",
    "main_amount",
    "transaction_amount",
    "amount_mmk",
    "amount",
    "paid_amount",
)
CEIR_TAX_PART_HEADERS = ("cd_mmk", "ct_mmk", "ait_mmk", "rf_mmk")
TXN_HEADERS = {
    "txn_id",
    "transaction_id",
    "txnid",
    "payment_id",
    "ref_no",
    "code",
}
SOURCE_HEADERS = {
    "source_type",
    "source",
    "type",
    "registration_type",
    "source_registration",
    "sourceregistration",
}
PAYMENT_DT_HEADERS = (
    "payment_dt",
    "payment_date",
    "paymentdate",
    "txn_date",
    "txndate",
    "transaction_date",
    "paid_at",
    "paid_date",
    "date",
)
CONFIRMED_AT_HEADERS = (
    "confirmed_at",
    "confirmation_date",
    "confirmed_date",
    "confirm_date",
    "ceir_confirmed_at",
    "confirmation_dt",
)

CEIR_APP_RE = re.compile(r"^MM-[A-Z0-9-]+$", re.IGNORECASE)
REQUEST_ID_RE = re.compile(r"^[A-Za-z0-9_-]{16,}$")


def normalize_header(value: Any) -> str:
    text = str(value if value is not None else "")
    text = text.replace("\n", " ").replace("\\", " ")
    text = text.strip().lower()
    text = re.sub(r"\s+", "_", text)
    text = re.sub(r"[^\w]", "", text)
    return text


def _to_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).replace(",", "").strip()
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _pick(mapping: dict[str, int], candidates: set[str] | tuple[str, ...]) -> int | None:
    for key in candidates:
        if key in mapping:
            return mapping[key]
    return None


def _cell_str(row: list[Any], col: int | None, limit: int) -> str | None:
    if col is None or col >= len(row) or row[col] is None:
        return None
    text = str(row[col]).strip()
    if not text or text.lower() == "none":
        return None
    return text[:limit]


def _parse_datetime(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        if value.tzinfo is None:
            return value.replace(tzinfo=timezone.utc)
        return value
    if isinstance(value, (int, float)):
        # Excel serial date (days since 1899-12-30).
        try:
            from datetime import timedelta

            base = datetime(1899, 12, 30, tzinfo=timezone.utc)
            return base + timedelta(days=float(value))
        except (OverflowError, ValueError):
            return None

    text = str(value).strip()
    if not text or text.lower() == "none":
        return None

    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M:%S%z",
        "%Y-%m-%d",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y/%m/%d",
    ):
        try:
            parsed = datetime.strptime(text.replace("Z", "+0000"), fmt)
            if parsed.tzinfo is None:
                parsed = parsed.replace(tzinfo=timezone.utc)
            return parsed
        except ValueError:
            continue
    return None


def _cell_datetime(row: list[Any], col: int | None) -> datetime | None:
    if col is None or col >= len(row):
        return None
    return _parse_datetime(row[col])


def _looks_like_header_row(mapped: dict[str, int]) -> bool:
    has_app = any(k in mapped for k in APP_ID_HEADERS)
    has_request = any(k in mapped for k in REQUEST_HEADERS)
    has_amount = _pick(mapped, AMOUNT_HEADER_PRIORITY + CEIR_TAX_PART_HEADERS) is not None
    has_status = _pick(mapped, STATUS_HEADERS) is not None
    if has_app and (has_request or has_status or has_amount):
        return True
    if has_request and (has_amount or has_status):
        return True
    return False


def _row_amount(row: list[Any], header_map: dict[str, int]) -> Decimal | None:
    # CEIR tax components first when present.
    parts: list[Decimal] = []
    for key in CEIR_TAX_PART_HEADERS:
        col = header_map.get(key)
        if col is None or col >= len(row):
            continue
        value = _to_decimal(row[col])
        if value is not None:
            parts.append(value)
    if parts:
        return sum(parts, Decimal("0"))

    for key in AMOUNT_HEADER_PRIORITY:
        col = header_map.get(key)
        if col is None or col >= len(row):
            continue
        value = _to_decimal(row[col])
        if value is not None:
            return value
    return None


def _rows_from_matrix(matrix: list[list[Any]]) -> list[dict[str, Any]]:
    if not matrix:
        return []

    header_row_index = -1
    header_map: dict[str, int] = {}
    for r, row in enumerate(matrix[:40]):
        mapped = {
            normalize_header(cell): i
            for i, cell in enumerate(row)
            if normalize_header(cell)
        }
        if _looks_like_header_row(mapped):
            header_row_index = r
            header_map = mapped
            break

    if header_row_index < 0 or not header_map:
        return []

    app_col = _pick(header_map, APP_ID_HEADERS)
    request_col = _pick(header_map, REQUEST_HEADERS)
    status_col = _pick(header_map, STATUS_HEADERS)
    txn_col = _pick(header_map, TXN_HEADERS)
    source_col = _pick(header_map, SOURCE_HEADERS)
    payment_dt_col = _pick(header_map, PAYMENT_DT_HEADERS)
    confirmed_at_col = _pick(header_map, CONFIRMED_AT_HEADERS)
    is_ceir_sheet = app_col is not None and any(
        k in header_map for k in ("application_id", "applicationid")
    )

    parsed: list[dict[str, Any]] = []
    for row in matrix[header_row_index + 1 :]:
        if not isinstance(row, list):
            row = list(row)

        app_id = _cell_str(row, app_col, 128)
        request_id = _cell_str(row, request_col, 128)

        if is_ceir_sheet:
            # CEIR rows must have a real Application ID.
            if not app_id or not CEIR_APP_RE.match(app_id):
                continue
        else:
            # Payment Hub: identity is request_id.
            if not app_id and request_id:
                app_id = request_id
            if not app_id or not (
                CEIR_APP_RE.match(app_id) or REQUEST_ID_RE.match(app_id)
            ):
                continue

        txn_id = _cell_str(row, txn_col, 128)
        payment_dt = _cell_datetime(row, payment_dt_col)
        confirmed_at = _cell_datetime(row, confirmed_at_col)
        # Hub sheets often only have one date column — reuse for payment_dt.
        if payment_dt is None and not is_ceir_sheet:
            payment_dt = confirmed_at
        if confirmed_at is None and is_ceir_sheet:
            confirmed_at = payment_dt

        parsed.append(
            {
                "application_id": app_id[:128],
                "request_id": (request_id[:128] if request_id else None),
                "status": _cell_str(row, status_col, 256),
                "tax_amount_mmk": _row_amount(row, header_map),
                "txn_id": (txn_id[:128] if txn_id else None),
                "source_type": _cell_str(row, source_col, 128),
                "payment_dt": payment_dt,
                "confirmed_at": confirmed_at,
            }
        )
    return parsed


def parse_csv_bytes(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    matrix = [list(row) for row in reader]
    return _rows_from_matrix(matrix)


def _matrix_from_calamine(data: bytes) -> list[list[Any]]:
    from python_calamine import CalamineWorkbook

    workbook = CalamineWorkbook.from_filelike(io.BytesIO(data))
    sheet = workbook.get_sheet_by_index(0)
    rows = sheet.to_python()
    return [list(row) if not isinstance(row, list) else row for row in rows]


def _matrix_from_openpyxl(data: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    for kwargs in (
        {"read_only": True, "data_only": True},
        {"read_only": False, "data_only": True},
        {"read_only": False, "data_only": False},
    ):
        try:
            workbook = load_workbook(filename=io.BytesIO(data), **kwargs)
            try:
                sheet = workbook.active
                return [list(row) for row in sheet.iter_rows(values_only=True)]
            finally:
                workbook.close()
        except Exception:
            continue
    raise ValueError(
        "Unable to read Excel workbook. Try re-saving as .xlsx or export as CSV."
    )


def parse_excel_bytes(data: bytes) -> list[dict[str, Any]]:
    errors: list[str] = []

    try:
        matrix = _matrix_from_calamine(data)
        rows = _rows_from_matrix(matrix)
        if rows:
            return rows
    except Exception as exc:
        errors.append(f"calamine: {exc}")

    try:
        matrix = _matrix_from_openpyxl(data)
        return _rows_from_matrix(matrix)
    except Exception as exc:
        errors.append(f"openpyxl: {exc}")

    detail = "; ".join(errors) if errors else "unknown parse error"
    raise ValueError(
        "Unable to read Excel workbook. Re-save as .xlsx or upload CSV. "
        f"({detail})"
    )


def parse_upload(filename: str, data: bytes) -> list[dict[str, Any]]:
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("File exceeds 25 MB limit")

    lower = filename.lower()
    if lower.endswith(".csv"):
        return parse_csv_bytes(data)
    if lower.endswith(".xlsx") or lower.endswith(".xls"):
        return parse_excel_bytes(data)
    raise ValueError("Only .csv, .xlsx, or .xls files are allowed")
